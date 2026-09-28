#!/usr/bin/env python3
"""Translate Armenian (hy) banking texts into Russian (ru) and/or English (en) with open-source models.

Engines:
  nllb-1.3b  facebook/nllb-200-distilled-1.3B (default, CC-BY-NC-4.0)
  nllb-600m  facebook/nllb-200-distilled-600M (smaller and faster, CC-BY-NC-4.0)
  opus       Helsinki-NLP/opus-mt-hy-{ru,en} (CC-BY-4.0 / Apache-2.0); hy->en pivots through Russian if needed

After translating, every paragraph is checked against the source: numbers (amounts, rates,
dates, account numbers) that are missing from the translation, and banking terms from
glossary.tsv whose expected translation is missing. These are printed as [review] lines
on stderr.
"""
import argparse
import os
import re
import sys
from collections import Counter

SRC = "hy"
TARGETS = ("ru", "en")
NAMES = {"hy": "Armenian", "ru": "Russian", "en": "English"}
NLLB_CODES = {"hy": "hye_Armn", "ru": "rus_Cyrl", "en": "eng_Latn"}
NLLB_MODELS = {
    "nllb-1.3b": "facebook/nllb-200-distilled-1.3B",
    "nllb-600m": "facebook/nllb-200-distilled-600M",
}
GLOSSARY_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "glossary.tsv")

# Sentence ends: Latin punctuation and the Armenian full stop (U+0589 "։").
# The Armenian question/exclamation marks sit inside the word, so "։" and "." end the sentence.
# A "." between digits (1.5, 15.03.2025) is not a sentence end because it is not followed by space.
# Many people type a plain ":" instead of "։", so ":" right after an Armenian letter also ends a sentence.
SENTENCE_END = re.compile(r"(?:(?<=[.!?։…])|(?<=[Ա-Ֆա-և]:))\s+")

# Numbers with thousands separators (1 000 000, 1,000,000) and decimals (12,5 / 12.5).
NUMBER = re.compile(r"\d{1,3}(?:[   .,]\d{3})+(?:[.,]\d+)?|\d+(?:[.,]\d+)?")
WORD = re.compile(r"\w+")


def armenian_share(text):
    """Fraction of letters in the text that are Armenian."""
    letters = [ch for ch in text if ch.isalpha()]
    armenian = [ch for ch in letters if "Ա" <= ch <= "֏" or "ﬓ" <= ch <= "ﬗ"]
    return len(armenian) / len(letters) if letters else 0.0


def split_sentences(paragraph, max_chars=400):
    """Split a paragraph into sentences, then cut any sentence that is still too long."""
    out = []
    for sentence in SENTENCE_END.split(paragraph.strip()):
        while len(sentence) > max_chars:
            cut = sentence.rfind(" ", 0, max_chars)
            cut = cut if cut > 0 else max_chars
            out.append(sentence[:cut])
            sentence = sentence[cut:].lstrip()
        if sentence:
            out.append(sentence)
    return out


# ---------------------------------------------------------------- checks

def numbers(text):
    """Digits of every number in the text, separators removed, so 1 000,50 and 1,000.50 compare equal."""
    return Counter(re.sub(r"\D", "", m) for m in NUMBER.findall(text))


def words(text):
    return [w.casefold().replace("ё", "е") for w in WORD.findall(text)]


def stem(word):
    """Crude stem for inflected words: drop up to two final letters, keep at least four."""
    return word if len(word) <= 4 else word[: max(4, len(word) - 2)]


def load_glossary(path=GLOSSARY_PATH):
    entries = []
    with open(path, encoding="utf-8") as f:
        rows = [line.rstrip("\n").split("\t") for line in f if line.strip() and not line.startswith("#")]
    header = rows[0]
    for row in rows[1:]:
        cells = dict(zip(header, row))
        entry = {lang: [v.strip() for v in cells[lang].split(";") if v.strip()] for lang in header}
        entries.append(entry)
    return entries


def term_matches(text_words, i, term, stemmed=True):
    """Number of words matched if `term` occurs at position i, else 0.

    Armenian glossary entries list their inflected forms explicitly, so they are matched
    unstemmed (otherwise "պարտք" would match "պարտավոր"); Russian and English are stemmed.
    """
    term_words = words(term)
    if i + len(term_words) > len(text_words):
        return 0
    for k, tw in enumerate(term_words):
        if not text_words[i + k].startswith(stem(tw) if stemmed else tw):
            return 0
    return len(term_words)


def glossary_hits(source, glossary):
    """Glossary entries found in the Armenian source, longest match first, no overlaps."""
    src_words = words(source)
    used = [False] * len(src_words)
    variants = sorted(
        ((variant, entry) for entry in glossary for variant in entry["hy"]),
        key=lambda ve: len(ve[0]),
        reverse=True,
    )
    hits = []
    for variant, entry in variants:
        for i in range(len(src_words)):
            n = term_matches(src_words, i, variant, stemmed=False)
            if n and not any(used[i : i + n]):
                used[i : i + n] = [True] * n
                if entry not in hits:
                    hits.append(entry)
    return hits


def contains_term(text, term):
    text_words = words(text)
    return any(term_matches(text_words, i, term) for i in range(len(text_words)))


def review_notes(source, translation, tgt, glossary):
    notes = []
    missing = numbers(source) - numbers(translation)
    if missing:
        notes.append("numbers missing or changed: " + ", ".join(sorted(missing.elements())))
    for entry in glossary_hits(source, glossary):
        if not any(contains_term(translation, t) for t in entry[tgt]):
            notes.append(f"term «{entry['hy'][0]}» should be «{entry[tgt][0]}», not found")
    return notes


# ---------------------------------------------------------------- engines

class NllbEngine:
    def __init__(self, model_name, device, beams):
        from transformers import AutoModelForSeq2SeqLM, AutoTokenizer

        self.tokenizer = AutoTokenizer.from_pretrained(model_name, src_lang=NLLB_CODES[SRC])
        self.model = AutoModelForSeq2SeqLM.from_pretrained(model_name).to(device)
        self.device = device
        self.beams = beams

    def translate(self, sentences, src, tgt):
        self.tokenizer.src_lang = NLLB_CODES[src]
        batch = self.tokenizer(sentences, return_tensors="pt", padding=True, truncation=True).to(self.device)
        out = self.model.generate(
            **batch,
            forced_bos_token_id=self.tokenizer.convert_tokens_to_ids(NLLB_CODES[tgt]),
            num_beams=self.beams,
            max_new_tokens=512,
        )
        return self.tokenizer.batch_decode(out, skip_special_tokens=True)


class OpusEngine:
    def __init__(self, device, beams):
        self.device = device
        self.beams = beams
        self.cache = {}

    def _load(self, src, tgt):
        key = (src, tgt)
        if key not in self.cache:
            from transformers import MarianMTModel, MarianTokenizer

            name = f"Helsinki-NLP/opus-mt-{src}-{tgt}"
            try:
                tok = MarianTokenizer.from_pretrained(name)
                model = MarianMTModel.from_pretrained(name).to(self.device)
                self.cache[key] = (tok, model)
            except OSError:
                self.cache[key] = None
        return self.cache[key]

    def _direct(self, sentences, pair):
        tok, model = pair
        batch = tok(sentences, return_tensors="pt", padding=True, truncation=True).to(self.device)
        out = model.generate(**batch, num_beams=self.beams, max_new_tokens=512)
        return tok.batch_decode(out, skip_special_tokens=True)

    def translate(self, sentences, src, tgt):
        pair = self._load(src, tgt)
        if pair:
            return self._direct(sentences, pair)
        # No direct hy->en model: pivot through Russian (hy->ru and ru->en both exist).
        first, second = self._load(src, "ru"), self._load("ru", tgt)
        if not (first and second):
            sys.exit(f"No OPUS-MT model for {src}->{tgt}, directly or via Russian. Use --engine nllb-1.3b.")
        print(f"[opus] no direct {src}->{tgt} model, pivoting through Russian", file=sys.stderr)
        return self._direct(self._direct(sentences, first), second)


# ---------------------------------------------------------------- translation

class Translator:
    def __init__(self, engine, glossary):
        self.engine = engine
        self.glossary = glossary
        self.notes = []

    def paragraph(self, text, tgt, where, batch_size=16):
        sentences = split_sentences(text.replace("\n", " "))
        translated = []
        for i in range(0, len(sentences), batch_size):
            translated.extend(self.engine.translate(sentences[i : i + batch_size], SRC, tgt))
        out = " ".join(translated)
        for note in review_notes(text, out, tgt, self.glossary):
            self.notes.append(f"[review] {NAMES[tgt]}, {where}: {note}")
        return out

    def text(self, text, tgt):
        """Translate plain text paragraph by paragraph so blank-line layout is kept."""
        result, n = [], 0
        for part in re.split(r"(\n\s*\n)", text):
            if part.strip():
                n += 1
                result.append(self.paragraph(part, tgt, f"paragraph {n}"))
            else:
                result.append(part)
        return "".join(result)

    def docx(self, path_in, path_out, tgt):
        """Translate every Armenian paragraph of a .docx (body, tables, headers, footers) in place."""
        import docx

        doc = docx.Document(path_in)
        counter = {"n": 0}

        def do_paragraphs(paragraphs, label):
            for p in paragraphs:
                if armenian_share(p.text) == 0:
                    continue
                counter["n"] += 1
                out = self.paragraph(p.text, tgt, f"{label} paragraph {counter['n']}")
                if p.runs:
                    # Keep the formatting of the first run; inline formatting changes inside the paragraph are lost.
                    p.runs[0].text = out
                    for run in p.runs[1:]:
                        run.text = ""
                else:
                    p.text = out

        def do_tables(tables, label):
            for table in tables:
                for row in table.rows:
                    for cell in row.cells:
                        do_paragraphs(cell.paragraphs, label)
                        do_tables(cell.tables, label)

        do_paragraphs(doc.paragraphs, "body")
        do_tables(doc.tables, "table")
        for section in doc.sections:
            for part, label in ((section.header, "header"), (section.footer, "footer")):
                do_paragraphs(part.paragraphs, label)
                do_tables(part.tables, label)
        doc.save(path_out)


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("text", nargs="*", help="Armenian text to translate (default: --input or stdin)")
    p.add_argument("--to", dest="tgt", default="all", choices=TARGETS + ("all",),
                   help="ru, en, or all for both (default: all)")
    p.add_argument("--engine", default="nllb-1.3b", choices=tuple(NLLB_MODELS) + ("opus",))
    p.add_argument("--input", help="read Armenian text from this .txt or .docx file")
    p.add_argument("--output", help="output file; for .docx with --to all, a directory (default: next to the input)")
    p.add_argument("--device", default=None, help="cpu or cuda (default: cuda if available)")
    p.add_argument("--beams", type=int, default=4)
    args = p.parse_args()

    is_docx = bool(args.input) and args.input.lower().endswith(".docx")
    if is_docx:
        import docx

        doc = docx.Document(args.input)
        text = "\n".join(p.text for p in doc.paragraphs)
        text += "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    elif args.text:
        text = " ".join(args.text)
    elif args.input:
        with open(args.input, encoding="utf-8") as f:
            text = f.read()
    else:
        text = sys.stdin.read()
    if not text.strip():
        sys.exit("Nothing to translate.")
    if armenian_share(text) < 0.5:
        sys.exit("The input does not look like Armenian. This skill translates from Armenian only.")

    targets = list(TARGETS) if args.tgt == "all" else [args.tgt]

    import torch

    device = args.device or ("cuda" if torch.cuda.is_available() else "cpu")
    if args.engine == "opus":
        engine = OpusEngine(device, args.beams)
    else:
        engine = NllbEngine(NLLB_MODELS[args.engine], device, args.beams)
    translator = Translator(engine, load_glossary())

    if is_docx:
        base = os.path.splitext(os.path.basename(args.input))[0]
        for tgt in targets:
            if args.output and len(targets) == 1 and not os.path.isdir(args.output):
                out_path = args.output
            else:
                out_dir = args.output or os.path.dirname(os.path.abspath(args.input))
                os.makedirs(out_dir, exist_ok=True)
                out_path = os.path.join(out_dir, f"{base}.{tgt}.docx")
            translator.docx(args.input, out_path, tgt)
            print(f"{NAMES[tgt]}: {out_path}")
    else:
        blocks = []
        for tgt in targets:
            out = translator.text(text, tgt)
            blocks.append(f"[{NAMES[tgt]}]\n{out}" if len(targets) > 1 else out)
        result = "\n\n".join(blocks)
        if args.output:
            with open(args.output, "w", encoding="utf-8") as f:
                f.write(result + "\n")
        else:
            print(result)

    for note in translator.notes:
        print(note, file=sys.stderr)
    if not translator.notes:
        print("[review] no missing numbers or glossary terms found", file=sys.stderr)


if __name__ == "__main__":
    main()
